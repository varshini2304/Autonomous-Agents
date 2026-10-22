from datetime import datetime
from pathlib import Path
from uuid import uuid4
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx(output_dir: Path, title: str, document_type: str, objective: str, assumptions: list[str], content: dict, outcomes: list[str]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.75); section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.85); section.right_margin = Inches(.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(45, 55, 65)
    for style_name, size, color in (("Title", 25, "17365D"), ("Heading 1", 16, "17365D"), ("Heading 2", 12, "2878A5")):
        style = doc.styles[style_name]; style.font.name = "Aptos Display"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(document_type.upper()).bold = True
    doc.add_heading(title, 0)
    p = doc.add_paragraph(f"Generated {datetime.now().astimezone().strftime('%B %d, %Y at %H:%M %Z')}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Objective", level=1); doc.add_paragraph(objective)
    if assumptions:
        doc.add_heading("Assumptions and constraints", level=1)
        for item in assumptions: doc.add_paragraph(item, style="List Bullet")
    for heading, body in content.items():
        doc.add_heading(heading, level=1)
        if isinstance(body, list):
            for item in body: doc.add_paragraph(str(item), style="List Bullet")
        elif isinstance(body, dict):
            for subheading, detail in body.items():
                doc.add_heading(str(subheading), level=2)
                if isinstance(detail, list):
                    for item in detail: doc.add_paragraph(str(item), style="List Bullet")
                else: doc.add_paragraph(str(detail))
        else: doc.add_paragraph(str(body))
    if outcomes:
        doc.add_heading("Execution outcomes", level=1)
        for item in outcomes: doc.add_paragraph(item, style="List Bullet")
    path = output_dir / f"{uuid4()}.docx"
    doc.save(path)
    return path
